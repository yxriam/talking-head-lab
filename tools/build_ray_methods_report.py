"""Build a source-grounded report answering Ray's four methodology questions."""
from pathlib import Path
import json
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output' / 'reports'
OUT.mkdir(parents=True, exist_ok=True)
NAME = 'facebook_voice_portrait_model_report_for_ray'
doc = Document()
md = []
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = Inches(.65)
sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.7)
sec.right_margin = Inches(.7)
for name, size in [('Normal',10.5),('Title',23),('Subtitle',11),('Heading 1',15),('Heading 2',11.5)]:
    st = doc.styles[name]
    st.font.name = 'Arial'
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0,0,0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.08
    if name.startswith('Heading'):
        st.font.bold = True
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.keep_with_next = True
    # The bundled Word template includes a decorative title paragraph border.
    # Remove it from styles so the document uses a plain black report title.
    if st._element.pPr is not None:
        for border in list(st._element.pPr.findall(qn('w:pBdr'))):
            st._element.pPr.remove(border)
    if name == 'Normal':
        st.paragraph_format.widow_control = True
header = sec.header.paragraphs[0]
header.text = 'Facebook voice and portrait experiment'
header.runs[0].font.size = Pt(8)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run('Prepared for Ray  |  8 October 2026  |  ')
r.font.size = Pt(8)
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
footer._p.append(fld)

def p(text, style=None):
    obj = doc.add_paragraph(text, style=style)
    md.append(text + '\n')
    return obj

def h(text, level=1):
    doc.add_heading(text, level)
    md.append('#'*(level+1) + ' ' + text + '\n')

def page():
    doc.add_page_break()

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for c,w in zip(t.columns,widths):
        c.width = Inches(w)
    for ri, values in enumerate([headers]+rows):
        row = t.rows[0] if ri == 0 else t.add_row()
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        if ri == 0:
            pr.append(OxmlElement('w:tblHeader'))
        for ci,(cell,text,width) in enumerate(zip(row.cells,values,widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ['top','bottom','left','right']:
                el = OxmlElement('w:'+edge)
                for key,val in [('val','single'),('sz','4'),('color','D9D9D9')]:
                    el.set(qn('w:'+key),val)
                borders.append(el)
            cp.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                el = OxmlElement('w:'+edge)
                el.set(qn('w:w'),'90')
                el.set(qn('w:type'),'dxa')
                margins.append(el)
            cp.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'),'E8EDF2' if ri == 0 else ('F7F8FA' if ri%2 == 0 else 'FFFFFF'))
            cp.append(shade)
            para=cell.paragraphs[0]
            para.paragraph_format.space_after=Pt(2)
            para.paragraph_format.space_before=Pt(2)
            para.paragraph_format.line_spacing=1.02
            para.alignment=WD_ALIGN_PARAGRAPH.CENTER if len(str(text))<20 and ci>0 else WD_ALIGN_PARAGRAPH.LEFT
            rr=para.add_run(str(text))
            rr.font.size=Pt(9)
            rr.bold=ri==0
    md.append('| ' + ' | '.join(headers) + ' |\n' + '| ' + ' | '.join(['---']*len(headers)) + ' |\n' + '\n'.join('| '+' | '.join(str(value).replace('\n','<br>') for value in row)+' |' for row in rows)+'\n')
    spacer=doc.add_paragraph()
    spacer.paragraph_format.space_after=Pt(1)
    spacer.paragraph_format.space_before=Pt(1)
    spacer.paragraph_format.line_spacing=Pt(2)

def linkline(label, url):
    para=doc.add_paragraph()
    para.paragraph_format.space_after=Pt(4)
    rid=para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    el=OxmlElement('w:hyperlink'); el.set(qn('r:id'),rid)
    run=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
    size=OxmlElement('w:sz'); size.set(qn('w:val'),'18'); rp.append(size)
    color=OxmlElement('w:color'); color.set(qn('w:val'),'174A73'); rp.append(color)
    run.append(rp); txt=OxmlElement('w:t'); txt.text=label; run.append(txt); el.append(run); para._p.append(el)
    md.append(f'[{label}]({url})\n')

doc.add_paragraph('Facebook Voice and Portrait Model Evaluation', style='Title')
md.append('# Facebook Voice and Portrait Model Evaluation\n')
p('Local PC results with a separate server history appendix', 'Subtitle')
p('8 October 2026')
p('The main report covers experiments on the local PC only. SadTalker 512 was the documented choice for the local Ray demonstration, while JoyVASA provided faster previews. The earlier rented-server workflow and the later hosted API route are described separately in the final appendix. Neither contributes numerical results to the local comparison tables.')
p('Local dataset R1 is the Ray 4-second pilot. Local dataset R2 is the fourteen-input benchmark. Local deliverable R3 is the 51-second Ray demonstration. Each is reported separately with its own inputs and settings. This report uses saved project evidence; no inference or Facebook collection was rerun.')
h('Voice sample collection',2)
p('For the historical Ray Facebook sample, I played the accessible video in the browser, routed its playback audio through VB-CABLE and recorded it with FFmpeg. The package preserves ray_video_vbcable.wav and a loudness-normalised MP3 copy. This was a recording of browser playback, rather than a direct download of the original Facebook audio track. [E1]')
p('The local Ray acceptance reference, ray_prompt_24k.wav, is a 20-second mono 24 kHz file. This reference is distinct from the approximately 2.4-second driving clips in R2 and the 4-second clip in R1. Historical audio collection describes where a source recording came from; it does not establish where a later model ran. Server-specific audio preparation is documented only in the appendix. [E1, E4, E6]')
h('Picture collection and preparation',2)
p('The retained Facebook image workflow used Playwright with a logged-in Chrome session, inspected page and network image URLs, and saved image files delivered by Facebook. The Ray image manifest records downloaded JPEGs, file sizes and hashes. It also states that screenshots were removed. The documented collection method is therefore downloading visible image files, rather than simply copying and pasting a screenshot. [E3]')
p('The Ray animation experiment uses acceptance/ray/ray.png. Its preparation script detects the largest face, makes a face-centred crop and resizes it to 768 × 768, without generating a replacement face. The reviewed records do not explicitly map ray.png to one particular Facebook download; that exact image provenance remains to be confirmed. [E4]')
h('How the text voice and picture are combined',2)
p('In the local Ray demonstration, written script + reference voice → Chatterbox speech audio; saved portrait + generated speech audio → SadTalker → MP4. The voice sample supplies the vocal reference, the text supplies the new words, and the portrait supplies the visual identity. This local sequence is evidenced by the acceptance scripts and manifests. [E4]')

page()
h('Models deployed on the local PC')
p('Ray’s question refers to LLMs, but the pipeline contains several different model types. The speech and portrait models have different inputs and roles. Only the text-processing component below is used as a text LLM.')
table(['Model or component','Role and inputs','Recorded execution'],[
    ['Chatterbox','Speech synthesis and voice cloning from reference audio and text','Local GPU execution recorded'],
    ['SadTalker','Talking-head video from a portrait and driving audio','Downloaded and run on the PC through Ubuntu in WSL'],
    ['EchoMimic V1','Audio-driven portrait animation','Downloaded and tested on the local PC'],
    ['JoyVASA','Audio-driven facial expression and head motion','Downloaded and tested on the local PC'],
    ['EchoMimic V3 Flash','Portrait video generation with image, audio and a scene prompt','Downloaded and tested locally using memory offloading'],
    ['Qwen3 1.7B Q8_0','Text LLM for scene selection and account analysis','Local llama.cpp with GPU support; not given Ray’s voice or picture'],
], [1.5,2.8,2.57])
p('The local deployment records identify an NVIDIA RTX 5070 Ti Laptop GPU with approximately 12 GB of VRAM and Ubuntu 22.04 under WSL. A browser page on localhost is a local interface; it does not imply that the model runs in the cloud. [E5]')
h('Local inputs and model roles',2)
p('The Ray acceptance script begins, “Hello, I am Ray, and I want to recommend Alice to your company.” The local Chatterbox script supplies this text and ray_prompt_24k.wav to the speech model. It records the generated speech before the local video model is run. This is the local text example; the different earlier server script appears in the appendix. [E4]')
p('The fourteen-input benchmark R2 feeds prepared driving recordings directly to each portrait model. It therefore evaluates the video stage, independently of the local long-script speech generation in R3. The model inventory describes local deployment; it does not imply that every listed model took the same script, voice reference and portrait in one end-to-end experiment. [E4, E6]')
h('Separation from remote work',2)
p('The earlier server stage used Chatterbox and SadTalker on a rented Linux GPU server. The hosted TokenHub HumanActor stage is a separate API route. Their workflows and source records are kept in the final appendix, without assigning local hardware, local timings or local identity scores to either remote stage.')

page()
h('Local R1 comparison using the Ray pilot')
p('The most directly relevant comparison uses the same Ray portrait and the same 4-second audio clip. SadTalker’s selected full-image configuration is compared with EchoMimic V1 and JoyVASA. The different output sizes and rendering methods are retained because they are part of the tested configurations. [E4]')
table(['Model','Output','Time in seconds','Identity mean','Identity P05'],[
    ['SadTalker 512 full pose 0','768 × 768\n25 fps','64.7','0.9561','0.9186'],
    ['EchoMimic V1 512','512 × 512\n24 fps','185.1','0.9524','0.9305'],
    ['JoyVASA','512 × 512\n25 fps','22.9','0.8365','0.7671'],
],[1.6,1.35,1.05,1.4,1.47])
p('Identity mean is the average ArcFace cosine similarity between the source face and sampled generated faces. P05 is the low fifth percentile, which exposes weaker frames. Higher scores indicate stronger similarity under this measurement; they are not percentages of identity accuracy.')
p('SadTalker was the preferred configuration for the Ray deliverable because its documented inspection combined strong identity retention with stable background and restrained motion. EchoMimic V1 retained identity well, including the highest P05 in this pilot, but required about 2.9 times as long and the project review noted visible mouth and face deformation. JoyVASA took about one third of SadTalker’s time and produced more mouth movement, but its lower identity scores show a clear trade-off. These visual observations are from the existing acceptance review, rather than a new blinded assessment. [E4]')
h('Local R2 comparison across fourteen inputs',2)
p('The wider benchmark combines seven portraits with two driving recordings, including the saved Facebook recording, for fourteen shared input combinations. Each prepared audio clip is approximately 2.4 seconds. It tests audio-driven animation directly; it does not compare how well different systems clone a voice or read a new script. All three models completed fourteen out of fourteen attempts. [E6]')
agg=json.loads((ROOT/'benchmark-results/2026-09-24-multi-sample-v2/aggregate-results.json').read_text(encoding='utf-8'))['models']
rows=[]
for key,label in [('sadtalker','SadTalker'),('echomimic_v1','EchoMimic V1'),('joyvasa','JoyVASA')]:
    a=agg[key]
    rows.append([label,f"{a['wall_seconds_mean']:.1f} ± {a['wall_seconds_std']:.1f}",f"{a['identity_mean_mean']:.4f}",f"{a['lse_d_mean']:.3f}",f"{a['peak_gpu_memory_mb_mean']:.0f}"])
table(['Model','Mean time in seconds','Identity mean','SyncNet LSE D','Peak GPU MB'],rows,[1.55,1.65,1.3,1.25,1.12])
p('Time variation is the population standard deviation across inputs. Lower LSE-D indicates better audio–visual synchronisation under SyncNet. These timings cover the model stage, excluding shared text, speech and background preparation. The older benchmark uses different configurations from the Ray 512 pilot, so the two tables should not be treated as one timing series.')
h('Additional EchoMimic V3 Flash result',2)
p('V3 Flash also completed fourteen out of fourteen benchmark cases. Its mean identity score was 0.8414. Two independent cold-start runs averaged 535.0 seconds, while twelve runs sharing a loaded model averaged 235.1 seconds each. Its resource report records approximately 13.3 GB of process memory. These measurements support keeping V3 as an experimental local option; they do not establish it as the best model for Ray’s portrait. [E6]')

page()
h('Local R3 demonstration and interpretation')
p('The saved natural-background Ray demonstration was produced with SadTalker 512 and is 51.52 seconds long at 768 × 768 and 25 fps. The recorded video generation stage took 667.5 seconds, or about 11.1 minutes. Its generated speech is 51.48 seconds long; Chatterbox speech generation took 155.3 seconds. These are separate measured stages and do not constitute a complete application latency measurement. [E4]')
p('Chatterbox generated the long script in four semantic chunks, joined with short pauses. SadTalker then rendered the combined driving audio in one continuous video inference. The saved quality check reports no black-frame segments and one silence interval longer than one second. The project review interprets that interval as a natural sentence pause. The continuous video was not assembled by stitching separately generated video segments. [E4]')
h('What the comparison supports',2)
p('For the documented Ray setup, SadTalker is the strongest supported choice for a presentation demonstration where identity retention and a stable background matter. JoyVASA is a useful fast preview option when more movement is desirable and some identity variation is acceptable. EchoMimic V1 is a useful diffusion-model comparison, although it did not provide a clear speed or observed visual advantage in this pilot.')
p('The records support comparison of the tested talking-head configurations. They do not support a three-model LLM comparison, a three-model voice-cloning comparison, or a claim that every model used the identical original Facebook picture and spoken script. Chatterbox is the documented speech system; Qwen processes text only. The exact mapping of the Ray source portrait, and matched evaluation of the cloud output, remain evidence gaps.')
p('The clips and sample set are small, model settings and output sizes differ, and neither face similarity nor SyncNet fully captures perceived naturalness. Results should be described as project measurements rather than general performance claims. This report has verified saved records and files, without rerunning inference or conducting a new listening test.')
h('Four sentences for a presentation',2)
p('1. I obtained the historical Facebook voice sample by recording browser playback through VB-CABLE with FFmpeg, and then prepared a trimmed mono WAV reference for Chatterbox.')
p('2. The Facebook collection saved image files delivered by the page through a Playwright browser workflow, while the Ray experiment prepared a face-centred crop from its saved source portrait.')
p('3. On my local PC, I used Chatterbox to generate speech from reference audio and text, and tested SadTalker, EchoMimic V1, JoyVASA and EchoMimic V3 Flash; the previous server experiments are documented separately.')
p('4. In the same-input Ray pilot, SadTalker offered the preferred overall balance, JoyVASA was fastest with more identity variation, and EchoMimic V1 preserved identity well but took substantially longer.')

page()
h('中文摘要与证据来源')
p('这份报告回应 Ray 提出的四项要求：Facebook 声音如何取得、图片如何保存、实际使用了哪些模型，以及三个模型的简要比较。历史 Ray 声音通过 VB-CABLE 将浏览器播放声音送入 FFmpeg 录制；图片采集记录显示通过 Playwright 获取页面提供的图片文件，并非仅复制粘贴截图。Ray 成片使用本地 ray.png 裁剪后的肖像，但现有记录尚未明确它对应哪一张 Facebook 下载图片。')
p('正文只报告本机实验：Chatterbox 根据参考声音和文字生成语音；SadTalker、EchoMimic V1、JoyVASA 和 EchoMimic V3 Flash 有本机视频测试记录。Qwen3 1.7B 是处理文字的本地 LLM。早期租用服务器上的 Chatterbox 与 SadTalker 流程，以及后来的 TokenHub HumanActor 云端路线，分别放在末尾附录，未将其数据放入本机对照表。')
p('本机 R1 是同一 Ray 肖像和四秒音频的对照：SadTalker 耗时 64.7 秒，身份均值 0.9561；EchoMimic V1 耗时 185.1 秒，身份均值 0.9524；JoyVASA 耗时 22.9 秒，身份均值 0.8365。本机 R2 是七张照片与两段约 2.4 秒音频组成的十四组实验。本机 R3 是 51.52 秒成片。三组输入、参数和目的不同，分别陈述，不能合并统计；这些数据均不代表服务器性能。')
p('已有 SadTalker 自然背景成片为 51.52 秒。云端输出也已保存，但缺少同条件指标、完整计时和费用记录，因此没有虚构云端排名或把它并入本机定量对照。报告整理的是已有实验，没有重新运行模型。')
h('Local results and collection evidence',2)
evidence=[
    ('E1  Historical Facebook voice package','ray_hunt_facebook_package/README.md'),
    ('E3  Original Facebook image download manifest','facebook-scam/data/accounts/01_ray_hunt/original_media/original_media_manifest.json'),
    ('E4  Ray pilot and completed demonstration','acceptance/ray/README.md'),
    ('E5  Local deployment records','local-media/README.md'),
    ('E6  Fourteen input benchmark results','benchmark-results/2026-09-24-multi-sample-v2/aggregate-results.json'),
]
for label,rel in evidence:
    linkline(label,(ROOT/rel).as_uri())
p('E1 and E3 establish collection history, not compute location. E4 contains the local R1 pilot, local R3 voice manifest and final performance logs. E6 contains local R2 results. E5 is supported by INSTALL-STATUS.md. The broad local-media README also discusses cloud integration; those passages are cited only in the separate API appendix. The historical image downloader is preserved in local-media/legacy-code-20261008.zip.')
h('Official model references',2)
for label,url in [
    ('Chatterbox speech synthesis and reference audio','https://github.com/resemble-ai/chatterbox'),
    ('SadTalker portrait and audio animation','https://github.com/OpenTalker/SadTalker'),
    ('EchoMimic audio driven portrait animation','https://github.com/antgroup/echomimic'),
    ('JoyVASA portrait animation','https://github.com/jdh-algo/JoyVASA'),
    ('EchoMimic V3 and Flash','https://github.com/antgroup/echomimic_v3'),
]:
    linkline(label,url)
p('Official references explain model roles. All numerical results in this report are from local project records, not vendor benchmarks. References checked on 8 October 2026; current upstream releases may differ from the project checkpoints.')

page()
h('Separate history of remote execution')
p('This appendix is separate from local datasets R1, R2 and R3. It describes earlier remote workflows using their own records. Local comparison values are not reused as server or hosted API results, and no local and remote scores are pooled.')
h('Earlier rented server workflow',2)
p('The historical WebUI documentation describes Chatterbox and SadTalker installed on a rented Linux GPU server. The Windows PC accessed the server through an SSH tunnel. The historical application and speech scripts use remote /root paths, unlike the /opt/media-models and /mnt/d/project paths in the local WSL acceptance scripts. These are two deployment stages, even though they share model names. [S1, S2]')
p('The server WebUI accepts a face image, reference audio or video, and text. Its FFmpeg step selects a reference segment, converts it to mono 24 kHz WAV, and applies filtering and loudness normalisation. Chatterbox generates the new speech; SadTalker renders the portrait video. Its image input supports upload and clipboard, with full-image or crop preparation. This describes the server implementation, not proof that a specific source photograph was acquired by copy and paste. [S1, S2]')
p('One earlier server speech script uses the reference filename ray2_british_prompt_24s.wav and begins, “This is an AI generated demo made with Ray’s consent.” The filename identifies that script’s reference; it does not replace the measured 20-second reference used in the later local acceptance experiment. [S2]')
p('The reviewed server documentation and scripts establish the pipeline, but do not supply a matched three-model server benchmark with verified hardware, timing and identity measurements. No such figures are inferred from the local experiments. Any quantitative claims in an earlier server report must retain that report’s own input, configuration, hardware and source attribution before being added here.')
h('Later hosted API route',2)
p('TokenHub HumanActor is a separate hosted service rather than the earlier rented-server installation. The integration uploads image and audio temporarily to Tencent COS, submits signed URLs to the API and downloads the result. The saved file acceptance/ray/ray_alice_recommendation_tokenhub_30s.mp4 records an existing output. [C1]')
p('There is no reviewed matched quality evaluation, full timing or cost record for that API output. It is therefore not given a numerical rank against the local models or the earlier rented server. Upload, network and service dependencies belong to this API route; local VRAM and runtime figures do not describe it.')
h('Remote evidence sources',2)
for label,rel in [
    ('S1  Earlier rented server WebUI documentation','TALKING_HEAD_WEBUI_README.md'),
    ('S2  Historical server application','talking_head_webui_app.py'),
    ('S2  Earlier server Ray speech script','generate_ray_british_disclosed_recommendation_audio.py'),
    ('C1  Hosted API adapter','local-media/tokenhub.py'),
]:
    linkline(label,(ROOT/rel).as_uri())

doc.core_properties.title='Facebook Voice and Portrait Model Evaluation'
doc.core_properties.subject='Report responding to Ray on collection methods and tested models'
doc.core_properties.author='CV Project'
doc.core_properties.keywords='Facebook, Chatterbox, SadTalker, EchoMimic, JoyVASA, Ray'
doc.save(OUT/(NAME+'.docx'))
(OUT/(NAME+'.md')).write_text('\n'.join(md),encoding='utf-8')
print(OUT/(NAME+'.docx'))
print(OUT/(NAME+'.md'))
