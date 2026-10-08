from __future__ import annotations

import json
import math
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\project\NZ\cv")
BENCH = ROOT / "benchmark-results" / "2026-09-23-equal-input"
OUT_DOCX = ROOT / "output" / "docx"
OUT_PDF = ROOT / "output" / "pdf"
ASSETS = ROOT / "output" / "report-assets"
DOCX_PATH = OUT_DOCX / "audio-driven-portrait-evaluation-report-zh.docx"

NAVY = "17324D"
TEAL = "2C6E6F"
MINT = "EAF3F1"
PALE = "F4F7F8"
GOLD = "D59B45"
RED = "B65050"
GRAY = "5E6B73"
LIGHT = "D9E2E7"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LIGHT, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, size=None, bold=None, color=None, east_asia="Microsoft YaHei"):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_para(doc, text="", style=None, bold_prefix=None, keep=False):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_run_font(p.add_run(item))
        p.paragraph_format.space_after = Pt(3)


def add_numbered(doc, items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.28)
        p.paragraph_format.first_line_indent = Inches(-.28)
        set_run_font(p.add_run(f"{index}.  {item}"))
        p.paragraph_format.space_after = Pt(3)


def add_table(doc, headers, rows, widths=None, font_size=8.5, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(h))
        set_run_font(r, font_size, True, WHITE)
        if widths:
            cell.width = widths[i]
    for row_data in rows:
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            set_cell_border(cell)
            if len(table.rows) % 2 == 0:
                set_cell_shading(cell, PALE)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_run_font(r, font_size, first_col_bold and i == 0)
            if widths:
                cell.width = widths[i]
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_callout(doc, title, text, color=TEAL):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_shading(cell, MINT if color == TEAL else "FFF4E5")
    set_cell_border(cell, color=color, size="8")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_run_font(r, 10, True, color)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run(text)
    set_run_font(r2, 9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, 8.5, color=GRAY)
    return p


def make_charts(data):
    ASSETS.mkdir(parents=True, exist_ok=True)
    labels = ["SadTalker", "EchoMimic V1", "JoyVASA", "EchoMimic V3 Flash"]
    colors = ["#2C6E6F", "#5D8AA8", "#D59B45", "#8B6F9C"]
    font_path = r"C:\Windows\Fonts\arial.ttf"
    bold_path = r"C:\Windows\Fonts\arialbd.ttf"
    font = ImageFont.truetype(font_path, 20)
    small = ImageFont.truetype(font_path, 17)
    bold = ImageFont.truetype(bold_path, 24)

    def single_chart(draw, box, title, values, formatter, minimum=0.0):
        x0, y0, x1, y1 = box
        draw.text((x0, y0), title, font=bold, fill="#17324D")
        plot_top, plot_bottom = y0 + 48, y1 - 68
        plot_left, plot_right = x0 + 25, x1 - 15
        draw.line((plot_left, plot_bottom, plot_right, plot_bottom), fill="#AFC0C8", width=2)
        vmax = max(values)
        span = max(vmax - minimum, vmax * .15, .01)
        bar_w = 72
        gap = (plot_right - plot_left) / len(values)
        for i, (value, label, color) in enumerate(zip(values, labels, colors)):
            cx = plot_left + gap * (i + .5)
            h = (value - minimum) / span * (plot_bottom - plot_top - 28)
            top = plot_bottom - max(h, 2)
            draw.rounded_rectangle((cx-bar_w/2, top, cx+bar_w/2, plot_bottom), radius=7, fill=color)
            val = formatter(value)
            tw = draw.textbbox((0, 0), val, font=small)[2]
            draw.text((cx-tw/2, top-24), val, font=small, fill="#17324D")
            lines = label.replace(" V3 Flash", " V3\nFlash").split("\n")
            for j, line in enumerate(lines):
                tw = draw.textbbox((0, 0), line, font=small)[2]
                draw.text((cx-tw/2, plot_bottom+10+j*18), line, font=small, fill="#5E6B73")

    img = Image.new("RGB", (1840, 650), "white")
    draw = ImageDraw.Draw(img)
    single_chart(draw, (45, 35, 900, 615), "Inference time (s, lower is better)", [d["wall_seconds"] for d in data], lambda v: f"{v:.1f}")
    single_chart(draw, (955, 35, 1795, 615), "Peak RAM RSS (GB, lower is better)", [d["max_rss_mb"]/1024 for d in data], lambda v: f"{v:.1f}")
    p1 = ASSETS / "runtime-resource.png"
    img.save(p1, quality=95)

    img = Image.new("RGB", (1840, 650), "white")
    draw = ImageDraw.Draw(img)
    single_chart(draw, (45, 35, 900, 615), "Identity similarity (higher is better)", [d["identity_mean"] for d in data], lambda v: f"{v:.3f}", minimum=.65)
    single_chart(draw, (955, 35, 1795, 615), "SyncNet LSE-C (higher is better)", [d["lse_c"] for d in data], lambda v: f"{v:.2f}")
    p2 = ASSETS / "identity-sync.png"
    img.save(p2, quality=95)
    return p1, p2


def setup_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(.62)
    sec.bottom_margin = Inches(.62)
    sec.left_margin = Inches(.72)
    sec.right_margin = Inches(.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(5)

    for name, size, color, space in [
        ("Title", 28, NAVY, 10),
        ("Heading 1", 18, NAVY, 8),
        ("Heading 2", 13, TEAL, 5),
        ("Heading 3", 10.5, NAVY, 3),
    ]:
        st = styles[name]
        heading_font = "SimHei" if name == "Heading 1" else "Microsoft YaHei"
        st.font.name = heading_font
        st._element.rPr.rFonts.set(qn("w:eastAsia"), heading_font)
        st.font.size = Pt(size)
        st.font.bold = name != "Heading 1"
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(space)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True

    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_run_font(header.add_run("音频驱动人像视频生成与评估实验报告"), 8.5, color=GRAY)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(footer.add_run("本地受控基准 · 2026-09-23    |    "), 8, color=GRAY)
        add_page_number(footer)
    return doc


def build_report():
    OUT_DOCX.mkdir(parents=True, exist_ok=True)
    OUT_PDF.mkdir(parents=True, exist_ok=True)
    data = json.loads((BENCH / "benchmark-summary.json").read_text(encoding="utf-8"))
    chart_runtime, chart_quality = make_charts(data)
    comparison = BENCH / "comparison-sheet.jpg"
    doc = setup_document()

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(72)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("音频驱动人像视频生成\n与评估实验报告")
    set_run_font(r, 28, True, NAVY)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("论文区域诊断实验复盘 × 本地四模型同输入基准")
    set_run_font(r, 14, color=TEAL)
    p2.paragraph_format.space_after = Pt(36)

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info = [
        ("报告范围", "Audio2Head、SadTalker、EchoMimic、JoyVASA、EchoMimic V3 Flash"),
        ("本地设备", "NVIDIA RTX 5070 Ti Laptop GPU 12 GB · WSL2 Ubuntu 22.04 · CUDA 12.8"),
        ("本地基准日期", "2026-09-23"),
        ("参考实验", "用户提供的 WACV 2027 匿名投稿稿件（11 页）"),
    ]
    for row, (k, v) in zip(table.rows, info):
        prevent_row_split(row)
        set_cell_shading(row.cells[0], MINT)
        for c in row.cells:
            set_cell_border(c, color=LIGHT)
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_run_font(row.cells[0].paragraphs[0].add_run(k), 9, True, TEAL)
        set_run_font(row.cells[1].paragraphs[0].add_run(v), 9)
    doc.add_paragraph().paragraph_format.space_after = Pt(22)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("用途：模型选型、部署决策、后续规模化评测设计"), 10, color=GRAY)
    doc.add_page_break()

    # Executive summary
    doc.add_heading("执行摘要", level=1)
    add_callout(doc, "核心结论", "在当前单张照片与约 2.4 秒语音的受控条件下，SadTalker 的身份保持和音画同步最稳；JoyVASA 的速度与嘴部运动幅度更突出；EchoMimic V1 的身份保持接近 SadTalker，但耗时和资源占用更高；EchoMimic V3 Flash 的画面清晰度最高，却在本样本上出现明显身份漂移、音画偏移和极慢推理。")
    add_bullets(doc, [
        "论文实验说明，仅靠全脸或整段视频的单一分数无法解释失败发生在何处。下颌区域会主导全脸外观误差，而眼部可能出现“几何稳定、外观漂移”的解耦。",
        "论文在 444 段 HDTF 诊断集上发现跨模型纹理误差长尾：前 4 个身份贡献约 64%–65% 的拉普拉斯纹理误差，前 10 个身份贡献约 85%–86%。",
        "本地四模型基准补齐了部署侧信息：推理时间、实时因子、显存、内存、身份相似度、SyncNet、嘴部运动和关键点稳定性。",
        "两组实验的输入规模、模型版本、硬件和指标实现不同，报告只比较趋势，不把绝对数值放在同一排名中。",
    ])
    doc.add_heading("决策建议", level=2)
    add_table(doc, ["使用目标", "当前推荐", "主要依据"], [
        ["身份稳定、口型同步、快速验证", "SadTalker", "身份均值 0.9638；LSE-C 4.641；约 21.6 秒"],
        ["更明显的头部/嘴部运动", "JoyVASA", "推理最快 20.2 秒；嘴部开合范围 0.2833"],
        ["EchoMimic 路线研究", "EchoMimic V1", "身份均值 0.9530，但耗时 59.4 秒、峰值显存 5.1 GB"],
        ["高细节探索", "EchoMimic V3 Flash", "清晰度 162.2；当前样本的身份与同步仍不合格"],
    ], widths=[Inches(1.7), Inches(1.6), Inches(3.5)], font_size=8.8, first_col_bold=True)

    doc.add_heading("目录", level=2)
    add_numbered(doc, [
        "研究背景与报告范围",
        "参考论文的实验设计与主要发现",
        "本地四模型同输入基准",
        "综合分析与模型选择",
        "限制、风险与下一轮实验设计",
        "指标释义与参考文献",
    ])
    doc.add_page_break()

    # Section 1
    doc.add_heading("1 研究背景与报告范围", level=1)
    add_para(doc, "本项目需要在本地或可控算力环境中完成“参考人像 + 驱动语音 → 人像视频”的生成，并同时评估身份保持、唇音同步、面部运动自然度、画面稳定性和部署成本。此前网页已接入多条生成路径，本报告聚焦可复现实验结果与选型依据。")
    doc.add_heading("1.1 为什么需要多维评估", level=2)
    add_para(doc, "音频驱动人像视频不是单一图像生成问题。输出可能在全局清晰度上很好，却出现眼部漂移、嘴形不自然、身份变化或音画错位。论文稿件进一步证明，全脸平均分会被面积较大的区域主导，因此不能替代区域级诊断。")
    add_table(doc, ["评价维度", "本报告使用的指标", "回答的问题"], [
        ["身份", "ArcFace 类嵌入余弦相似度：均值、P05、最低值、标准差", "输出是否还是同一个人，漂移是否集中在少数帧"],
        ["音画同步", "SyncNet LSE-D、LSE-C、最优偏移帧", "嘴部运动与音频节奏是否对齐"],
        ["嘴形与稳定性", "开合范围、嘴部抖动、嘴宽 CV、不对称、全脸关键点抖动", "运动是否足够、是否稳定、是否存在变形"],
        ["效率", "墙钟时间、RTF、峰值显存、最大 RSS", "能否在当前设备长期部署"],
        ["论文区域诊断", "五区域外观/几何、十纹理块、I3D-FVD、身份崩溃", "全局分数背后的具体失败位置和类型"],
    ], widths=[Inches(1.2), Inches(3.2), Inches(2.5)], font_size=8.4)
    doc.add_heading("1.2 数据来源", level=2)
    add_bullets(doc, [
        "参考论文：Region-Level Diagnostics for Talking-Head Evaluation: Area Bias, Texture Concentration, and Identity Stress Cases，用户提供的 WACV 2027 匿名投稿稿件。",
        "本地实验：D:\\project\\NZ\\cv\\benchmark-results\\2026-09-23-equal-input，包含原始请求、模型日志、资源采样、视频输出和逐模型指标。",
        "本地输入：同一张 768×768 肖像与同一段约 2.4 秒单人语音；共享的场景理解、背景生成和人物抠图阶段不计入各模型推理时间。",
    ])

    # Section 2
    doc.add_heading("2 参考论文的实验设计与主要发现", level=1)
    doc.add_heading("2.1 三层数据设计", level=2)
    add_table(doc, ["层级", "规模", "用途", "注意事项"], [
        ["HDTF 诊断子集", "444 段；37 个身份 × 每身份 12 段", "完整区域外观、几何、纹理、身份和来源因素分析", "每段均用三个生成器评估"],
        ["HDTF 扩展验证", "960 段；80 个身份 × 每身份 12 段", "ArcFace、FaceNet、I3D-FVD、SyncNet 聚合验证", "固定随机种子 42"],
        ["CelebV-HQ 迁移集", "61 个三模型共同成功样本；35 个文件名定义的身份组", "跨数据集趋势验证", "从 200 个候选到 95 个清洗通过；受 EchoMimic 成功率筛选影响"],
    ], widths=[Inches(1.3), Inches(1.7), Inches(2.6), Inches(1.5)], font_size=8.2)
    add_para(doc, "论文比较 Audio2Head、SadTalker 和 EchoMimic，均使用作者默认检查点、单张源图和同身份语音。区域诊断均匀抽取 5 帧，I3D-FVD 抽取 16 帧；脸部先通过 FAN 68 点对齐到 512×512。")
    doc.add_heading("2.2 区域诊断协议", level=2)
    add_bullets(doc, [
        "五个语义区域：眼、鼻、口、下颌、脸颊；双侧区域取左右统计的平均。",
        "外观误差：源帧和生成帧对应区域掩膜并集上的平均 L1 差异。",
        "几何误差：区域关键点位移除以源帧瞳距，减少分辨率和人脸尺度影响。",
        "十个皮肤纹理块：额头、眼下、眼角、脸颊皮肤、鼻唇沟和下巴；使用拉普拉斯方差差与 LBP 直方图距离。",
        "身份：ArcFace CSIM，RetinaFace 阈值为 0.1；CSIM < 0.1 定义为身份崩溃。FaceNet 用作独立识别器复核。",
    ])
    add_callout(doc, "解释边界", "区域指标把生成帧与静态源图比较，因此口部和下颌的差异包含预期说话运动。论文主要把这些指标用于身份、纹理与异常诊断，而不把它们直接称为“口型准确率”。", color=GOLD)
    doc.add_heading("2.3 主要量化结果", level=2)
    add_table(doc, ["模型", "I3D-FVD ↓", "平均 CSIM（444 段）", "身份崩溃率"], [
        ["Audio2Head", "11.69", "0.370", "24.55%，95% CI [13.5, 36.7]"],
        ["SadTalker", "11.75", "0.940", "0%；单侧 95% 上界 <0.7%"],
        ["EchoMimic", "10.72", "0.931", "0%；单侧 95% 上界 <0.7%"],
    ], widths=[Inches(1.5), Inches(1.3), Inches(1.8), Inches(2.5)], font_size=8.5, first_col_bold=True)
    add_para(doc, "在 960 段 HDTF 扩展验证中，ArcFace CSIM 为 Audio2Head 0.255、SadTalker 0.941、EchoMimic 0.926；在 61 段 CelebV-HQ 迁移集上分别为 0.118、0.852、0.746，身份排名保持一致。")
    add_table(doc, ["区域", "Audio2Head 正常样本", "SadTalker", "EchoMimic"], [
        ["眼", "27.24 / 0.024", "18.91 / 0.022", "19.62 / 0.021"],
        ["鼻", "21.66 / 0.069", "13.92 / 0.040", "15.22 / 0.042"],
        ["口", "32.70 / 0.092", "26.99 / 0.060", "30.75 / 0.065"],
        ["下颌", "19.19 / 0.124", "12.22 / 0.073", "12.78 / 0.067"],
        ["脸颊", "17.86 / 0.100", "10.33 / 0.057", "10.63 / 0.058"],
        ["全脸", "19.68 / 0.079", "12.13 / 0.050", "12.66 / 0.050"],
    ], widths=[Inches(1.1), Inches(2.0), Inches(1.8), Inches(1.8)], font_size=8.2)
    add_caption(doc, "表 1 论文 444 段诊断集的区域统计：外观 L1 / 瞳距归一化几何位移。Audio2Head 仅统计 CSIM ≥ 0.5 的正常样本。")
    doc.add_heading("2.4 论文提出的诊断结论", level=2)
    add_bullets(doc, [
        "面积偏置：下颌外观误差与全脸外观误差几乎完全相关，三个生成器的 Pearson r 均大于 0.98；眼部相关性仅 0.399–0.657。",
        "眼部外观—几何解耦：眼部几何误差约为全脸的 0.3–0.4 倍，但眼部外观误差约为全脸的 1.4–1.6 倍。",
        "纹理误差长尾：前 4 个身份贡献约 64%–65% 的拉普拉斯纹理误差，前 10 个身份贡献约 85%–86%；模型间身份难度排序的 Spearman 相关系数均超过 0.98。",
        "来源因素：源帧 RMS 对比度与生成纹理误差相关最强，Audio2Head/SadTalker/EchoMimic 的 r 分别为 0.633/0.626/0.661，均 p < 0.001；亮度变化也显著相关。",
        "身份应力案例：Audio2Head 的崩溃在身份间高度集中，且眼、鼻、脸颊等非口部区域的几何误差同样上升，说明并非正常说话运动。FaceNet 对 ArcFace 定义的崩溃与正常组分离 AUC 为 1.000。",
        "多指标排名冲突：在 CelebV-HQ 上，Audio2Head 的 I3D-FVD 最低，SadTalker 的像素/感知与身份指标最好，EchoMimic 的 SyncNet 最好；单一指标不足以决定模型优劣。",
    ])
    # Section 3
    doc.add_heading("3 本地四模型同输入基准", level=1)
    doc.add_heading("3.1 环境与控制变量", level=2)
    add_table(doc, ["项目", "设置"], [
        ["硬件", "NVIDIA RTX 5070 Ti Laptop GPU，12 GB 显存"],
        ["环境", "Windows + WSL2 Ubuntu 22.04；CUDA 12.8"],
        ["输入图像", "同一张 768×768 原像素人物 / 新背景肖像"],
        ["输入音频", "同一段约 2.4 秒单人语音"],
        ["计时边界", "仅模型推理阶段；排除共享的 Qwen 场景分类、SDXL 背景生成和 U2Net 人物提取"],
        ["进程策略", "每个模型独立新进程；未清空操作系统文件缓存"],
        ["评测样本量", "每个模型 1 个同输入视频；身份/关键点按输出的 56–63 帧统计"],
    ], widths=[Inches(1.5), Inches(5.5)], font_size=8.7, first_col_bold=True)
    add_callout(doc, "可复现性说明", "本地结果适合比较同一工作站上的模型行为和部署成本，但不能当作跨人物总体准确率。单样本还不能给出置信区间，也不能替代论文中的 444/960 段规模化评测。", color=GOLD)

    doc.add_heading("3.2 速度与资源", level=2)
    add_table(doc, ["模型", "时间 s ↓", "RTF ↓", "峰值显存 MB ↓", "最大 RSS MB ↓", "输出"], [
        ["SadTalker", "21.6", "8.89", "2,154", "2,237", "256×256 / 25 fps"],
        ["EchoMimic V1", "59.4", "23.77", "5,194", "6,810", "384×384 / 24 fps"],
        ["JoyVASA", "20.2", "8.03", "2,494", "3,165", "512×448 / 25 fps"],
        ["EchoMimic V3 Flash", "269.7", "117.30", "2,272*", "7,009", "384×384 / 25 fps"],
    ], widths=[Inches(1.55), Inches(.85), Inches(.75), Inches(1.15), Inches(1.15), Inches(1.45)], font_size=8.2, first_col_bold=True)
    add_para(doc, "* EchoMimic V3 Flash 的低显存峰值来自 CPU offload；它同时具有最高内存占用和最长耗时，因此不能解释为整体资源更省。")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(chart_runtime), width=Inches(6.8))
    add_caption(doc, "图 1 本地四模型的推理时间与资源峰值。")

    doc.add_heading("3.3 身份保持与唇音同步", level=2)
    add_table(doc, ["模型", "身份均值 ↑", "P05 ↑", "最低 ↑", "LSE-D ↓", "LSE-C ↑", "偏移帧 →0"], [
        ["SadTalker", "0.9638", "0.9444", "0.9359", "9.640", "4.641", "−1"],
        ["EchoMimic V1", "0.9530", "0.9091", "0.9039", "11.921", "1.222", "−2"],
        ["JoyVASA", "0.8675", "0.8088", "0.7784", "11.263", "3.050", "−1"],
        ["EchoMimic V3 Flash", "0.7767", "0.7195", "0.6998", "12.853", "1.204", "−13"],
    ], widths=[Inches(1.55), Inches(.9), Inches(.75), Inches(.75), Inches(.75), Inches(.75), Inches(.95)], font_size=8.1, first_col_bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(chart_quality), width=Inches(6.8))
    add_caption(doc, "图 2 本地身份相似度与 SyncNet LSE-C。LSE-C 越高通常表示音画同步置信度越高。")

    doc.add_heading("3.4 嘴形、关键点与画面稳定性", level=2)
    add_table(doc, ["模型", "开合范围", "嘴抖 ↓", "嘴宽 CV ↓", "不对称 ↓", "全脸抖动 ↓", "几何异常率 ↓", "清晰度 ↑"], [
        ["SadTalker", "0.1112", "0.0133", "0.0192", "0.0021", "0.0090", "0%", "42.2"],
        ["EchoMimic V1", "0.1232", "0.0111", "0.0192", "0.0020", "0.0149", "10%", "20.5"],
        ["JoyVASA", "0.2833", "0.0214", "0.0623", "0.0058", "0.0107", "0%", "22.7"],
        ["EchoMimic V3 Flash", "0.1276", "0.0108", "0.0256", "0.0048", "0.0111", "0%", "162.2"],
    ], widths=[Inches(1.35), Inches(.8), Inches(.7), Inches(.75), Inches(.75), Inches(.9), Inches(.9), Inches(.75)], font_size=7.7, first_col_bold=True)
    add_para(doc, "嘴部开合范围反映动作幅度，不是越大越好；嘴抖、宽度变异和不对称需要结合动作幅度与视觉检查判断。四个输出的人脸检测率均为 100%。")
    if comparison.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(comparison), width=Inches(6.9))
        add_caption(doc, "图 3 相同输入下的逐模型关键帧对照。定量分数用于筛查，最终仍需结合连续视频观察身份漂移、眼部与嘴部自然度。")

    # Section 4
    doc.add_heading("4 综合分析与模型选择", level=1)
    doc.add_heading("4.1 本地结果与论文发现的相互印证", level=2)
    add_table(doc, ["现象", "论文证据", "本地证据", "解释"], [
        ["身份保持排序", "HDTF 中 SadTalker 0.940、EchoMimic 0.931", "SadTalker 0.9638、EchoMimic V1 0.9530", "在不同协议中，两者均表现为高身份保持；绝对值不可直接比较"],
        ["单指标会冲突", "EchoMimic I3D-FVD 最低，但 SadTalker 身份更高", "V3 清晰度最高但身份与同步最弱", "模型选择必须同时看身份、同步、稳定性和效率"],
        ["区域异常可能被均值隐藏", "眼部外观误差被全脸平均稀释", "V1 具有 10% 几何异常帧，但身份均值仍为 0.9530", "需保留低分位、异常率与逐帧可视化"],
        ["更强运动会带来代价", "论文提示口/颌差异包含预期运动", "JoyVASA 开合最大，同时嘴宽 CV、嘴抖和身份漂移也更大", "运动幅度与身份稳定存在任务相关权衡"],
    ], widths=[Inches(1.2), Inches(1.9), Inches(1.9), Inches(2.0)], font_size=7.9)
    doc.add_heading("4.2 逐模型判断", level=2)
    add_bullets(doc, [
        "SadTalker：本机综合最稳。身份均值和最低值最高，SyncNet 最优，显存约 2.1 GB。主要不足是分辨率低、口部动作幅度偏小，画面容易显得僵硬。",
        "EchoMimic V1：身份保持接近 SadTalker，但推理约慢 2.7 倍，显存约 5.1 GB；本样本还出现 10% 几何异常帧与较低 SyncNet 置信度。适合保留为扩散模型对照。",
        "JoyVASA：速度略优于 SadTalker，输出尺寸更大、动作更明显；身份均值下降到 0.8675，嘴部抖动和宽度变化最高。适合重视表现力且允许一定身份变化的预览。",
        "EchoMimic V3 Flash：细节清晰，但 CPU offload 使 RTF 达 117.3，音画偏移 −13 帧，身份均值最低。当前配置不应作为默认路径。",
    ])
    doc.add_heading("4.3 推荐部署策略", level=2)
    add_numbered(doc, [
        "默认生成采用 SadTalker，向用户暴露高清增强与动作强度设置，并在输出后执行身份和同步质量门控。",
        "将 JoyVASA 作为“表现力优先”选项；若身份均值低于 0.90 或最低帧低于 0.82，则回退 SadTalker 或提示重选照片。",
        "EchoMimic V1 作为研究与高质量候选，先优化推理参数、分辨率和同步后再升级为正式选项。",
        "暂时隐藏 EchoMimic V3 Flash 的默认入口，只在实验模式中保留；先解决时序偏移与身份漂移。",
        "网页结果页同时显示身份低分位、音画偏移和异常帧时间段，避免仅展示一个平均分。",
    ])

    # Section 5
    doc.add_heading("5 限制、风险与下一轮实验设计", level=1)
    doc.add_heading("5.1 当前报告的限制", level=2)
    add_bullets(doc, [
        "本地基准只有一个人物和一段短语音，不能推断跨性别、年龄、肤色、姿态、光照、遮挡和语言的总体性能。",
        "本地身份指标与论文 ArcFace 管线可能使用不同人脸检测、裁剪和权重，绝对数值不具备严格可比性。",
        "SyncNet 的 LSE-D/LSE-C 是音画对齐代理指标，不是带真实口型标注的唇形准确率；偏移帧还受视频帧率和裁剪质量影响。",
        "清晰度是拉普拉斯类高频指标，锐化、噪声和纹理伪影都可能抬高分数，不能单独代表视觉质量。",
        "独立进程减少了显存残留影响，但没有清空系统文件缓存；时延应理解为受控工作站测量，而不是严格冷启动。",
        "论文的 CelebV-HQ 比较只保留三个模型共同成功的 61 段，受 EchoMimic 完成率选择偏差影响。",
    ])
    doc.add_heading("5.2 下一轮规模化实验", level=2)
    add_table(doc, ["阶段", "建议规模", "新增指标", "通过标准"], [
        ["冒烟集", "10 人 × 2 音频", "身份、同步、失败率、耗时", "所有模型管线可完成；无系统性偏移"],
        ["开发集", "30–50 人 × 3 条音频", "P05/P01、区域眼/口诊断、眨眼与头姿", "身份和同步阈值按开发集校准"],
        ["验证集", "≥100 人，身份隔离", "I3D-FVD、人工 MOS、区域纹理长尾", "先冻结参数，再一次性报告"],
        ["压力集", "侧脸、遮挡、低照、眼镜、强表情", "异常帧时段、失败类型、回退率", "按失败类型设定可接受上限"],
    ], widths=[Inches(1.1), Inches(1.5), Inches(2.8), Inches(1.6)], font_size=8.2)
    add_para(doc, "建议采用身份级划分与身份级 bootstrap 置信区间，避免把同一人物的多段视频当作独立样本。应记录源帧亮度变化、RMS 对比度、脸占比和姿态，因为论文已显示这些因素可能与纹理难度相关。")
    doc.add_heading("5.3 建议的质量门控", level=2)
    add_bullets(doc, [
        "硬失败：无人脸、音频缺失、输出时长异常、解码失败或最优音画偏移超过预设上限。",
        "身份门控：同时使用均值、P05、最低值和标准差；若均值尚可但最低值骤降，应标记具体秒段而不是放行。",
        "区域门控：眼部和嘴部单独报告外观/几何异常，避免下颌面积掩盖局部问题。",
        "人工抽检：每个模型至少查看开头、中段、结尾和自动标记的异常帧，并按身份、口型、眼部、姿态、闪烁五项打分。",
    ])

    # Appendix
    doc.add_heading("6 指标释义与参考文献", level=1)
    doc.add_heading("6.1 指标方向", level=2)
    add_table(doc, ["指标", "方向", "含义与使用注意"], [
        ["身份 CSIM", "越高越好", "源人脸与生成帧的人脸嵌入余弦相似度；需记录检测器、阈值和裁剪方式"],
        ["LSE-D", "越低越好", "SyncNet 音画距离；不同实现和裁剪下不可盲目比较"],
        ["LSE-C", "越高越好", "SyncNet 同步置信度；应与最优偏移帧一起报告"],
        ["RTF", "越低越好", "推理时间 / 输出视频时长；RTF=1 表示实时"],
        ["I3D-FVD", "越低越好", "视频分布级时序真实性；不能定位具体失败区域"],
        ["P05", "越高越好", "逐帧指标的第 5 百分位，用于发现平均值掩盖的低质量尾部"],
        ["几何异常率", "越低越好", "关键点配置显著偏离正常范围的帧占比"],
    ], widths=[Inches(1.2), Inches(1.1), Inches(4.7)], font_size=8.4)
    doc.add_heading("6.2 参考文献", level=2)
    refs = [
        "用户提供稿件. Region-Level Diagnostics for Talking-Head Evaluation: Area Bias, Texture Concentration, and Identity Stress Cases. WACV 2027 anonymous submission.",
        "Wang S, Li L, Ding Y, et al. Audio2Head: Audio-driven One-shot Talking-head Generation with Natural Head Motion. IJCAI, 2021.",
        "Zhang W, Cun X, Wang X, et al. SadTalker: Learning Realistic 3D Motion Coefficients for Stylized Audio-Driven Single Image Talking Face Animation. CVPR, 2023.",
        "Chen Z, Cao J, Chen Z, et al. EchoMimic: Lifelike Audio-Driven Portrait Animations through Editable Landmark Conditions. arXiv:2407.08136, 2024.",
        "Chung J S, Zisserman A. Out of Time: Automated Lip Sync in the Wild. ACCV Workshops, 2016.",
        "Deng J, Guo J, Xue N, Zafeiriou S. ArcFace: Additive Angular Margin Loss for Deep Face Recognition. CVPR, 2019.",
        "Unterthiner T, van Steenkiste S, Kurach K, et al. Towards Accurate Generative Models of Video: A New Metric and Challenges. arXiv:1812.01717, 2018.",
    ]
    for i, ref in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.2)
        p.paragraph_format.first_line_indent = Inches(-.2)
        r = p.add_run(f"[{i}] {ref}")
        set_run_font(r, 8.5)

    doc.add_heading("6.3 复现文件", level=2)
    p = doc.add_paragraph()
    set_run_font(p.add_run("本地基准目录："), 9, True)
    set_run_font(p.add_run(str(BENCH)), 9)
    p2 = doc.add_paragraph()
    set_run_font(p2.add_run("SyncNet 官方实现："), 9, True)
    add_hyperlink(p2, "joonson/syncnet_python", "https://github.com/joonson/syncnet_python")
    add_para(doc, "报告中的本地数值来自 benchmark-summary.json、visual-summary.json、syncnet-summary.json 和逐模型 performance.json；模型输出与日志保留在各自子目录中。")

    # document metadata
    props = doc.core_properties
    props.title = "音频驱动人像视频生成与评估实验报告"
    props.subject = "论文区域诊断实验与本地四模型同输入基准"
    props.author = "CV Project"
    props.keywords = "Talking Head, SadTalker, EchoMimic, JoyVASA, SyncNet, ArcFace"
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_report()
