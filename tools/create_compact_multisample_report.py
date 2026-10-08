"""Create a compact Chinese report focused on multi-sample metrics and examples."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from create_experiment_report import (
    GRAY, NAVY, TEAL, WHITE, add_caption, add_hyperlink, add_page_number,
    add_para, add_table, set_run_font,
)


ROOT = Path(r"D:\project\NZ\cv")
BENCH = ROOT / "benchmark-results" / "2026-09-24-multi-sample-v2"
ASSETS = ROOT / "output" / "report-assets"
GALLERY = ROOT / "output" / "multi-sample-report" / "video-gallery.html"
OUT = ROOT / "output" / "docx" / "audio-driven-portrait-multisample-metrics-report-zh.docx"
MODELS = ("sadtalker", "echomimic_v1", "joyvasa", "echomimic_v3_flash")
LABELS = {"sadtalker": "SadTalker", "echomimic_v1": "EchoMimic V1", "joyvasa": "JoyVASA", "echomimic_v3_flash": "EchoMimic V3 Flash"}


def val(mean, std, digits=3, lower=False):
    if mean is None:
        return "—"
    arrow = "↓" if lower else "↑"
    return f"{mean:.{digits}f} ± {std:.{digits}f} {arrow}"


def case_metric_rows():
    visuals = json.loads((BENCH / "visual-results.json").read_text(encoding="utf-8"))
    vindex = {(x["case_id"], x["model"]): x for x in visuals if x.get("model") and "error" not in x}
    manifest = json.loads((BENCH / "manifest.json").read_text(encoding="utf-8"))
    winners = {"identity": {m: 0 for m in MODELS}, "sync": {m: 0 for m in MODELS}, "speed": {m: 0 for m in MODELS}}
    valid_identity = valid_sync = valid_speed = 0
    for item in manifest:
        cid = item["case_id"]
        scores = {m: vindex[(cid, m)]["identity_p05"] for m in MODELS if (cid, m) in vindex}
        if scores:
            winners["identity"][max(scores, key=scores.get)] += 1; valid_identity += 1
        sync = {}
        speed = {}
        for model in MODELS:
            s = BENCH / "cases" / cid / model / "syncnet-metrics.json"
            p = BENCH / "cases" / cid / model / "performance.json"
            if s.exists(): sync[model] = json.loads(s.read_text())["lse_d"]
            if p.exists():
                row = json.loads(p.read_text())
                if row.get("exit_code") == 0: speed[model] = row["wall_seconds"]
        if sync:
            winners["sync"][min(sync, key=sync.get)] += 1; valid_sync += 1
        if speed:
            winners["speed"][min(speed, key=speed.get)] += 1; valid_speed += 1
    return winners, (valid_identity, valid_sync, valid_speed)


def setup(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(.62); sec.bottom_margin = Inches(.58)
    sec.left_margin = Inches(.68); sec.right_margin = Inches(.68)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(5)
    for name, size, color in (("Title", 23, NAVY), ("Heading 1", 16, NAVY), ("Heading 2", 11.5, TEAL)):
        style = styles[name]
        style.font.name = "Aptos"; style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")
        style.font.size = Pt(size); style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = False if name == "Heading 1" else True
    header = sec.header.paragraphs[0]
    set_run_font(header.add_run("音频驱动人像 · 多样本基准"), 8.5, color=GRAY)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("CV Project  |  "), 8, color=GRAY); add_page_number(footer)


def image_page(doc, heading, image, caption):
    doc.add_heading(heading, level=1)
    doc.add_picture(str(image), width=Inches(7.0))
    add_caption(doc, caption)


def build():
    agg = json.loads((BENCH / "aggregate-results.json").read_text(encoding="utf-8"))
    winners, valid = case_metric_rows()
    doc = Document(); setup(doc)
    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("音频驱动人像多样本指标报告"); set_run_font(r, 23, False, NAVY, east_asia="SimHei")
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(sub.add_run("7 张照片 × 2 段语音 = 14 组同条件输入"), 11, color=TEAL)
    add_para(doc, "本报告只记录可复核指标、效果样本和跨样本共性。每个可运行模型使用完全相同的 14 组输入；失败结果保留并计入稳定性。输出片段约 2.2–2.5 秒，固定预处理和推理参数。")
    doc.add_heading("汇总指标", level=1)
    rows = []
    for model in MODELS:
        m = agg["models"][model]
        rows.append([
            LABELS[model], f'{m["successes"]}/{m["attempted"]}',
            val(m.get("identity_mean_mean"), m.get("identity_mean_std")),
            val(m.get("identity_p05_mean"), m.get("identity_p05_std")),
            val(m.get("lse_d_mean"), m.get("lse_d_std"), 2, True),
            val(m.get("lse_c_mean"), m.get("lse_c_std"), 2),
            val(m.get("wall_seconds_mean"), m.get("wall_seconds_std"), 1, True),
        ])
    add_table(doc, ["模型", "成功/尝试", "身份均值", "身份 P05", "LSE-D", "LSE-C", "耗时/s"], rows,
              widths=[Inches(1.15), Inches(.72), Inches(1.05), Inches(1.05), Inches(.95), Inches(.95), Inches(.95)], font_size=7.8, first_col_bold=True)
    add_para(doc, "方向：身份相似度与 LSE-C 越高越好；LSE-D 与耗时越低越好。± 后为跨样本总体标准差。身份 P05 是每段视频 11–12 个均匀时点的人脸相似度低 5% 分位，可揭示平均值掩盖的变形时点。V3 耗时栏采用两次独立冷启动实测；其余样本一次加载后连续运行，另列批量吞吐。")
    doc.add_heading("资源与稳定性", level=2)
    resource_rows = []
    for model in MODELS:
        m = agg["models"][model]
        resource_rows.append([LABELS[model], f'{m["success_rate"]*100:.0f}%' if m["success_rate"] is not None else "—",
                              val(m.get("max_rss_mb_mean"), m.get("max_rss_mb_std"), 0, True),
                              val(m.get("peak_gpu_memory_mb_mean"), m.get("peak_gpu_memory_mb_std"), 0, True),
                              val(m.get("face_detection_rate_mean"), m.get("face_detection_rate_std"))])
    add_table(doc, ["模型", "尝试内成功率", "内存/MB", "显存/MB", "人脸检出率"], resource_rows,
              widths=[Inches(1.4), Inches(1.05), Inches(1.55), Inches(1.55), Inches(1.45)], font_size=8.2, first_col_bold=True)
    doc.add_page_break()
    image_page(doc, "效果截图示例（1/2）", ASSETS / "multi-sample-contact-1.jpg", "图 1　四个人物样本的中间帧；所有列来自同一张输入照片与同一段音频。")
    doc.add_page_break()
    image_page(doc, "效果截图示例（2/2）", ASSETS / "multi-sample-contact-2.jpg", "图 2　其余三个人物样本；所有模型均完成 14 组生成。")
    doc.add_heading("跨样本共性", level=2)
    identity = sorted(winners["identity"].items(), key=lambda x: x[1], reverse=True)
    sync = sorted(winners["sync"].items(), key=lambda x: x[1], reverse=True)
    speed = sorted(winners["speed"].items(), key=lambda x: x[1], reverse=True)
    common = [
        ["低分位身份最佳", f"{LABELS[identity[0][0]]}：{identity[0][1]}/{valid[0]} 组"],
        ["音画距离最佳", f"{LABELS[sync[0][0]]}：{sync[0][1]}/{valid[1]} 组"],
        ["推理最快", f"{LABELS[speed[0][0]]}：{speed[0][1]}/{valid[2]} 组"],
        ["V3 Flash 稳定性", f'{agg["models"]["echomimic_v3_flash"]["successes"]}/{agg["models"]["echomimic_v3_flash"]["attempted"]} 组成功'],
        ["V3 批量吞吐", val(agg["models"]["echomimic_v3_flash"].get("warm_batch_wall_seconds_mean"), agg["models"]["echomimic_v3_flash"].get("warm_batch_wall_seconds_std"), 1, True) + " 秒/组（单次模型加载后）"],
    ]
    add_table(doc, ["观察", "跨样本计数"], common, widths=[Inches(2.1), Inches(4.9)], font_size=9, first_col_bold=True)
    doc.add_page_break()
    doc.add_heading("效果视频与复核路径", level=1)
    add_para(doc, "网页画廊包含全部 14 组输入下的逐模型输出，按样本并排播放；生成失败项同样显示。它用于复核截图无法表达的口型、眨眼、头动和帧间稳定性。")
    p = doc.add_paragraph(); set_run_font(p.add_run("视频画廊："), 10, True)
    add_hyperlink(p, str(GALLERY), GALLERY.as_uri())
    add_para(doc, f"完整输出、性能记录与日志：{BENCH}")
    doc.add_heading("公平性与解释边界", level=2)
    add_table(doc, ["控制项", "本次做法"], [
        ["输入一致", "同一批 7 张照片、2 段语音；每个模型使用相同 14 个组合"],
        ["样本保留", "不剔除低质量结果；失败计入成功率并保留日志"],
        ["统计稳定性", "每段取 11–12 个均匀时点；报告跨样本均值、总体标准差、P05 和逐样本胜出次数"],
        ["V3 执行方式", "2 组独立冷启动用于端到端耗时；其余 12 组固定同一权重实例连续推理，并单独记录每组耗时"],
        ["可复核性", "保留逐模型视频、性能 JSON、SyncNet 输出和模型日志"],
        ["适用范围", "14 组可用于本机工程比较；不足以代表所有人物、语言和拍摄条件"],
    ], widths=[Inches(1.35), Inches(5.65)], font_size=8.8, first_col_bold=True)
    add_para(doc, "结论应结合身份相似度、低分位、SyncNet 和实际视频观看；任何单一指标都不能完整代表主观自然度。扩大到更多身份时，应保持身份隔离并继续报告失败率与分布。")
    props = doc.core_properties
    props.title = "音频驱动人像多样本指标报告"; props.author = "CV Project"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT); print(OUT)


if __name__ == "__main__":
    build()
